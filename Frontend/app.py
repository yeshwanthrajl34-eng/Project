import streamlit as st
import requests
from io import BytesIO
from docx import Document
from fpdf import FPDF


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="LegalEase — AI Legal Workspace",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ============================================================
# SESSION STATE
# ============================================================

if "document" not in st.session_state:
    st.session_state.document = ""

if "generated" not in st.session_state:
    st.session_state.generated = False


# ============================================================
# PREMIUM UI CSS
# ============================================================

st.markdown(
    """
<style>

/* =========================================================
   IMPORT FONT
   ========================================================= */

@import url(
    'https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Playfair+Display:wght@600;700&display=swap'
);


/* =========================================================
   GLOBAL
   ========================================================= */

html, body, [class*="css"] {
    font-family: 'DM Sans', sans-serif;
}

.stApp {
    background:
        radial-gradient(
            circle at 10% 0%,
            rgba(99, 102, 241, 0.12),
            transparent 28%
        ),
        radial-gradient(
            circle at 90% 10%,
            rgba(6, 182, 212, 0.10),
            transparent 25%
        ),
        #f6f8fc;
}

.block-container {
    max-width: 1450px;
    padding-top: 2rem;
    padding-bottom: 3rem;
}


/* =========================================================
   SIDEBAR
   ========================================================= */

section[data-testid="stSidebar"] {
    background:
        linear-gradient(
            180deg,
            #0b1120 0%,
            #111827 45%,
            #171554 100%
        );

    border-right: 1px solid rgba(255,255,255,0.06);
}

section[data-testid="stSidebar"] * {
    color: #e5e7eb;
}

.sidebar-brand {
    text-align: center;
    padding: 18px 10px 28px;
}

.sidebar-logo {
    width: 62px;
    height: 62px;
    margin: auto;

    display: flex;
    align-items: center;
    justify-content: center;

    border-radius: 20px;

    background:
        linear-gradient(
            135deg,
            #6366f1,
            #8b5cf6,
            #06b6d4
        );

    box-shadow:
        0 12px 35px rgba(99,102,241,0.35);

    font-size: 31px;
}

.sidebar-title {
    font-family: 'Playfair Display', serif;
    font-size: 27px;
    font-weight: 700;
    color: white;
    margin-top: 13px;
}

.sidebar-subtitle {
    font-size: 11px;
    color: #94a3b8;
    letter-spacing: 0.5px;
    margin-top: 3px;
}

.sidebar-label {
    color: #64748b;
    font-size: 10px;
    font-weight: 700;
    letter-spacing: 1.5px;
    text-transform: uppercase;

    margin-top: 24px;
    margin-bottom: 10px;
}

.sidebar-item {
    padding: 10px 12px;
    margin-bottom: 6px;

    border-radius: 10px;

    background: rgba(255,255,255,0.035);

    color: #cbd5e1;

    font-size: 12px;
}

.sidebar-item-active {
    background:
        linear-gradient(
            90deg,
            rgba(99,102,241,0.25),
            rgba(139,92,246,0.10)
        );

    border: 1px solid rgba(129,140,248,0.18);

    color: white;
}


/* =========================================================
   HERO
   ========================================================= */

.hero {
    position: relative;
    overflow: hidden;

    padding: 42px 46px;

    border-radius: 28px;

    background:
        linear-gradient(
            120deg,
            #0b1120 0%,
            #172554 48%,
            #312e81 100%
        );

    box-shadow:
        0 25px 60px rgba(15,23,42,0.18);

    margin-bottom: 30px;
}

.hero::before {
    content: "";

    position: absolute;

    width: 320px;
    height: 320px;

    right: -100px;
    top: -130px;

    border-radius: 50%;

    background:
        radial-gradient(
            circle,
            rgba(34,211,238,0.20),
            transparent 68%
        );
}

.hero::after {
    content: "⚖";

    position: absolute;

    right: 55px;
    bottom: -25px;

    font-family: serif;

    font-size: 160px;

    color: rgba(255,255,255,0.035);
}

.hero-badge {
    display: inline-block;

    padding: 7px 13px;

    border-radius: 999px;

    background: rgba(255,255,255,0.08);

    border: 1px solid rgba(255,255,255,0.12);

    color: #c7d2fe;

    font-size: 10px;
    font-weight: 700;

    letter-spacing: 1.3px;

    margin-bottom: 16px;
}

.hero-title {
    position: relative;
    z-index: 2;

    font-family: 'Playfair Display', serif;

    font-size: clamp(35px, 4vw, 55px);

    line-height: 1.05;

    color: white;

    margin: 0;
}

.hero-title span {
    background:
        linear-gradient(
            90deg,
            #a5b4fc,
            #67e8f9
        );

    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
}

.hero-description {
    position: relative;
    z-index: 2;

    max-width: 720px;

    margin-top: 15px;

    color: #cbd5e1;

    font-size: 14px;

    line-height: 1.7;
}


/* =========================================================
   TOP STATUS
   ========================================================= */

.status-row {
    display: flex;
    align-items: center;
    gap: 8px;

    margin-bottom: 22px;

    font-size: 11px;
    color: #64748b;
}

.status-dot {
    width: 8px;
    height: 8px;

    border-radius: 50%;

    background: #22c55e;

    box-shadow:
        0 0 0 4px rgba(34,197,94,0.10);
}


/* =========================================================
   CARDS
   ========================================================= */

.card {
    background: rgba(255,255,255,0.92);

    border: 1px solid rgba(226,232,240,0.9);

    border-radius: 20px;

    padding: 24px;

    box-shadow:
        0 12px 35px rgba(15,23,42,0.055);

    backdrop-filter: blur(10px);

    margin-bottom: 20px;
}

.card-heading {
    display: flex;
    align-items: center;

    gap: 12px;

    margin-bottom: 20px;
}

.card-icon {
    width: 42px;
    height: 42px;

    display: flex;
    align-items: center;
    justify-content: center;

    border-radius: 13px;

    background:
        linear-gradient(
            135deg,
            #eef2ff,
            #ecfeff
        );

    border: 1px solid #e0e7ff;

    font-size: 20px;
}

.card-title {
    color: #111827;

    font-size: 18px;
    font-weight: 700;
}

.card-subtitle {
    color: #94a3b8;

    font-size: 11px;

    margin-top: 2px;
}


/* =========================================================
   INPUT LABELS
   ========================================================= */

label {
    font-weight: 600 !important;
    color: #334155 !important;
}


/* =========================================================
   INPUT BOXES
   ========================================================= */

div[data-baseweb="select"] > div {
    border-radius: 12px !important;

    border-color: #e2e8f0 !important;

    background: white !important;
}

textarea {
    border-radius: 12px !important;
}


/* =========================================================
   PRIMARY BUTTON
   ========================================================= */

.stButton > button {
    min-height: 48px;

    border-radius: 12px;

    font-weight: 700;

    border: none;

    transition:
        transform 0.2s ease,
        box-shadow 0.2s ease;
}

.stButton > button:hover {
    transform: translateY(-2px);

    box-shadow:
        0 10px 25px rgba(79,70,229,0.18);
}


/* =========================================================
   DOCUMENT HEADER
   ========================================================= */

.document-toolbar {
    display: flex;

    justify-content: space-between;
    align-items: center;

    padding: 11px 14px;

    background: #f8fafc;

    border:
        1px solid #e2e8f0;

    border-bottom: none;

    border-radius: 12px 12px 0 0;

    font-size: 11px;

    color: #64748b;
}

.ready-pill {
    display: inline-flex;

    align-items: center;

    gap: 5px;

    padding: 5px 10px;

    border-radius: 999px;

    background: #ecfdf5;

    color: #047857;

    font-weight: 700;
}


/* =========================================================
   STAT CARDS
   ========================================================= */

.stat {
    text-align: center;

    padding: 12px;

    border-radius: 13px;

    background: #f8fafc;

    border: 1px solid #e2e8f0;
}

.stat-number {
    color: #111827;

    font-size: 20px;

    font-weight: 800;
}

.stat-label {
    color: #94a3b8;

    font-size: 10px;

    margin-top: 2px;

    text-transform: uppercase;

    letter-spacing: 0.7px;
}


/* =========================================================
   EMPTY DOCUMENT
   ========================================================= */

.empty-state {
    text-align: center;

    padding: 100px 25px;

    border-radius: 18px;

    border: 1px dashed #cbd5e1;

    background:
        linear-gradient(
            145deg,
            #f8fafc,
            #eef2ff
        );
}

.empty-icon {
    font-size: 58px;

    margin-bottom: 15px;
}

.empty-title {
    color: #334155;

    font-size: 19px;

    font-weight: 700;
}

.empty-text {
    color: #94a3b8;

    font-size: 12px;

    margin-top: 7px;
}


/* =========================================================
   DOWNLOAD SECTION
   ========================================================= */

.export-heading {
    color: #111827;

    font-size: 16px;

    font-weight: 800;

    margin-top: 22px;

    margin-bottom: 4px;
}

.export-subtitle {
    color: #94a3b8;

    font-size: 11px;

    margin-bottom: 12px;
}


/* =========================================================
   DISCLAIMER
   ========================================================= */

.disclaimer {
    margin-top: 18px;

    padding: 13px 15px;

    border-radius: 12px;

    background: #fffbeb;

    border: 1px solid #fde68a;

    color: #92400e;

    font-size: 11px;

    line-height: 1.6;
}


/* =========================================================
   FOOTER
   ========================================================= */

.footer {
    text-align: center;

    margin-top: 35px;

    padding-top: 20px;

    border-top: 1px solid #e2e8f0;

    color: #94a3b8;

    font-size: 11px;
}

.footer strong {
    color: #6366f1;
}


/* =========================================================
   MOBILE
   ========================================================= */

@media (max-width: 900px) {

    .hero {
        padding: 30px 25px;
    }

    .hero-title {
        font-size: 35px;
    }

    .hero::after {
        display: none;
    }

}

</style>
""",
    unsafe_allow_html=True
)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.markdown(
        """
        <div class="sidebar-brand">

            <div class="sidebar-logo">
                ⚖️
            </div>

            <div class="sidebar-title">
                LegalEase
            </div>

            <div class="sidebar-subtitle">
                AI LEGAL WORKSPACE
            </div>

        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="sidebar-label">Workspace</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="sidebar-item sidebar-item-active">✦ Document Generator</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="sidebar-item">✎ Editable Workspace</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="sidebar-item">⇩ Document Export</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="sidebar-label">Document Types</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        """
        <div class="sidebar-item">🔐 Non-Disclosure Agreement</div>
        <div class="sidebar-item">💼 Employment Agreement</div>
        <div class="sidebar-item">⌂ Lease Agreement</div>
        <div class="sidebar-item">🤝 Service Agreement</div>
        <div class="sidebar-item">◈ Partnership Agreement</div>
        """,
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="sidebar-label">System</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        """
        <div class="sidebar-item">
            🟢 AI Engine Connected
        </div>

        <div class="sidebar-item">
            🔒 Local Workspace
        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown("---")

    st.caption(
        "LegalEase creates AI-assisted drafts. "
        "Always review documents before legal use."
    )


# ============================================================
# HERO
# ============================================================

st.markdown(
    """
    <div class="hero">

        <div class="hero-badge">
            ✦ INTELLIGENT LEGAL DOCUMENT WORKSPACE
        </div>

        <div class="hero-title">
            Legal documents,<br>
            <span>simplified by AI.</span>
        </div>

        <div class="hero-description">
            Describe your agreement, provide the key details,
            and LegalEase will create a structured legal draft
            that you can review, edit and export.
        </div>

    </div>
    """,
    unsafe_allow_html=True
)


# ============================================================
# STATUS
# ============================================================

st.markdown(
    """
    <div class="status-row">
        <span class="status-dot"></span>
        LegalEase AI workspace is ready
    </div>
    """,
    unsafe_allow_html=True
)


# ============================================================
# MAIN COLUMNS
# ============================================================

left, right = st.columns(
    [0.85, 1.5],
    gap="large"
)


# ============================================================
# LEFT — INPUT PANEL
# ============================================================

with left:

    st.markdown(
        '<div class="card">',
        unsafe_allow_html=True
    )

    st.markdown(
        """
        <div class="card-heading">

            <div class="card-icon">
                ✦
            </div>

            <div>
                <div class="card-title">
                    Create a Document
                </div>

                <div class="card-subtitle">
                    Provide the details for your agreement
                </div>
            </div>

        </div>
        """,
        unsafe_allow_html=True
    )

    document_type = st.selectbox(
        "Document type",

        [
            "Non-Disclosure Agreement",
            "Employment Agreement",
            "Lease Agreement",
            "Service Agreement",
            "Partnership Agreement",
            "Other"
        ]
    )

    parties = st.text_area(
        "Parties",

        placeholder=(
            "Example:\n"
            "ABC Technologies Pvt. Ltd.\n"
            "John Doe"
        ),

        height=105
    )

    dates = st.text_area(
        "Important dates",

        placeholder=(
            "Example:\n"
            "Effective Date: 6 October 2026\n"
            "Term: 12 months"
        ),

        height=95
    )

    terms = st.text_area(
        "Terms & requirements",

        placeholder=(
            "Describe the important terms of the agreement.\n\n"
            "Example: confidentiality obligations, payment terms, "
            "termination conditions, responsibilities, etc."
        ),

        height=185
    )

    st.markdown(
        "<div style='height:8px'></div>",
        unsafe_allow_html=True
    )

    generate = st.button(
        "✦  Generate with LegalEase AI",
        type="primary",
        use_container_width=True
    )

    clear = st.button(
        "↻  Clear Workspace",
        use_container_width=True
    )

    if clear:

        st.session_state.document = ""
        st.session_state.generated = False

        st.rerun()

    st.markdown(
        """
        <div class="disclaimer">

        <strong>Before you generate</strong><br>

        Provide accurate information about the parties,
        dates and important terms. LegalEase generates a
        draft for review and customization.

        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown(
        "</div>",
        unsafe_allow_html=True
    )


# ============================================================
# GENERATE
# ============================================================

if generate:

    if not parties.strip():

        st.warning(
            "Please enter the parties involved."
        )

    elif not terms.strip():

        st.warning(
            "Please enter the terms and requirements."
        )

    else:

        with st.spinner(
            "LegalEase AI is preparing your document..."
        ):

            try:

                response = requests.post(

                    "http://127.0.0.1:8000/generate",

                    json={
                        "document_type": document_type,
                        "parties": parties,
                        "dates": dates,
                        "terms": terms
                    },

                    timeout=180
                )

                if response.status_code == 200:

                    data = response.json()

                    generated_document = data.get(
                        "document",
                        ""
                    )

                    if generated_document:

                        st.session_state.document = (
                            generated_document
                        )

                        st.session_state.generated = True

                        st.success(
                            "Document generated successfully."
                        )

                    else:

                        st.error(
                            "The AI returned an empty document."
                        )

                else:

                    st.error(
                        f"Backend error: {response.text}"
                    )

            except requests.exceptions.ConnectionError:

                st.error(
                    "Cannot connect to LegalEase AI."
                )

                st.info(
                    "Make sure FastAPI is running with:"
                )

                st.code(
                    "uvicorn backend.main:app --reload"
                )

            except requests.exceptions.Timeout:

                st.error(
                    "The AI request timed out. "
                    "Please try generating again."
                )

            except Exception as error:

                st.error(
                    f"Unexpected error: {error}"
                )


# ============================================================
# RIGHT — DOCUMENT WORKSPACE
# ============================================================

with right:

    st.markdown(
        '<div class="card">',
        unsafe_allow_html=True
    )

    st.markdown(
        """
        <div class="card-heading">

            <div class="card-icon">
                ◈
            </div>

            <div>
                <div class="card-title">
                    Document Workspace
                </div>

                <div class="card-subtitle">
                    Review and edit your AI-generated draft
                </div>
            </div>

        </div>
        """,
        unsafe_allow_html=True
    )


    # ========================================================
    # DOCUMENT EXISTS
    # ========================================================

    if st.session_state.document:

        document = st.session_state.document

        words = len(document.split())

        characters = len(document)

        paragraphs = len(
            [
                p
                for p in document.split("\n")
                if p.strip()
            ]
        )


        # ----------------------------------------------------
        # STATS
        # ----------------------------------------------------

        stat1, stat2, stat3 = st.columns(3)

        with stat1:

            st.markdown(
                f"""
                <div class="stat">

                    <div class="stat-number">
                        {words}
                    </div>

                    <div class="stat-label">
                        Words
                    </div>

                </div>
                """,
                unsafe_allow_html=True
            )

        with stat2:

            st.markdown(
                f"""
                <div class="stat">

                    <div class="stat-number">
                        {paragraphs}
                    </div>

                    <div class="stat-label">
                        Sections
                    </div>

                </div>
                """,
                unsafe_allow_html=True
            )

        with stat3:

            st.markdown(
                f"""
                <div class="stat">

                    <div class="stat-number">
                        {characters}
                    </div>

                    <div class="stat-label">
                        Characters
                    </div>

                </div>
                """,
                unsafe_allow_html=True
            )


        st.markdown(
            "<div style='height:16px'></div>",
            unsafe_allow_html=True
        )


        # ----------------------------------------------------
        # EDITOR TOOLBAR
        # ----------------------------------------------------

        st.markdown(
            """
            <div class="document-toolbar">

                <span>
                    ✎ Editable legal draft
                </span>

                <span class="ready-pill">
                    ● AI Draft Ready
                </span>

            </div>
            """,
            unsafe_allow_html=True
        )


        # ----------------------------------------------------
        # EDITOR
        # ----------------------------------------------------

        edited_document = st.text_area(

            "Document editor",

            value=document,

            height=560,

            label_visibility="collapsed"
        )

        st.session_state.document = edited_document


        # ====================================================
        # EXPORT
        # ====================================================

        st.markdown(
            """
            <div class="export-heading">
                Export your document
            </div>

            <div class="export-subtitle">
                Download the edited version in your preferred format.
            </div>
            """,
            unsafe_allow_html=True
        )


        # ----------------------------------------------------
        # TXT
        # ----------------------------------------------------

        txt_data = edited_document.encode(
            "utf-8"
        )


        # ----------------------------------------------------
        # DOCX
        # ----------------------------------------------------

        doc = Document()

        title = doc.add_heading(
            "LegalEase",
            level=1
        )

        title.alignment = 1

        subtitle = doc.add_paragraph(
            "AI-Assisted Legal Document"
        )

        subtitle.alignment = 1

        doc.add_paragraph("")

        for paragraph in edited_document.split("\n"):

            if paragraph.strip():

                doc.add_paragraph(
                    paragraph.strip()
                )

            else:

                doc.add_paragraph("")


        docx_buffer = BytesIO()

        doc.save(
            docx_buffer
        )

        docx_buffer.seek(0)


        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        pdf = FPDF()

        pdf.set_auto_page_break(
            auto=True,
            margin=18
        )

        pdf.add_page()


        pdf.set_font(
            "Arial",
            "B",
            20
        )

        pdf.cell(
            0,
            12,
            "LegalEase",
            ln=True,
            align="C"
        )


        pdf.set_font(
            "Arial",
            "",
            9
        )

        pdf.cell(
            0,
            7,
            "AI-Assisted Legal Document",
            ln=True,
            align="C"
        )


        pdf.ln(8)


        pdf.set_font(
            "Arial",
            "",
            11
        )


        for line in edited_document.split("\n"):

            clean_line = (
                line
                .replace("–", "-")
                .replace("—", "-")
                .replace("’", "'")
                .replace("“", '"')
                .replace("”", '"')
                .replace("•", "-")
            )

            if clean_line.strip():

                pdf.multi_cell(
                    0,
                    7,
                    clean_line
                )

                pdf.ln(1)

            else:

                pdf.ln(4)


        pdf_output = bytes(
            pdf.output()
        )


        # ----------------------------------------------------
        # DOWNLOAD BUTTONS
        # ----------------------------------------------------

        txt_col, docx_col, pdf_col = st.columns(3)


        with txt_col:

            st.download_button(

                "↓  TXT",

                data=txt_data,

                file_name="LegalEase_Document.txt",

                mime="text/plain",

                use_container_width=True
            )


        with docx_col:

            st.download_button(

                "↓  DOCX",

                data=docx_buffer,

                file_name="LegalEase_Document.docx",

                mime=(
                    "application/"
                    "vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),

                use_container_width=True
            )


        with pdf_col:

            st.download_button(

                "↓  PDF",

                data=pdf_output,

                file_name="LegalEase_Document.pdf",

                mime="application/pdf",

                use_container_width=True
            )


        # ----------------------------------------------------
        # DISCLAIMER
        # ----------------------------------------------------

        st.markdown(
            """
            <div class="disclaimer">

                ⚠️ <strong>Legal review recommended.</strong>

                LegalEase provides AI-assisted document drafts
                for educational and drafting purposes. Review
                and customize the document with a qualified
                legal professional before relying on it.

            </div>
            """,
            unsafe_allow_html=True
        )


    # ========================================================
    # EMPTY STATE
    # ========================================================

    else:

        st.markdown(
            """
            <div class="empty-state">

                <div class="empty-icon">
                    📄
                </div>

                <div class="empty-title">
                    Your document will appear here
                </div>

                <div class="empty-text">
                    Complete the document details on the left,
                    then let LegalEase AI create your draft.
                </div>

            </div>
            """,
            unsafe_allow_html=True
        )


    st.markdown(
        "</div>",
        unsafe_allow_html=True
    )


# ============================================================
# FOOTER
# ============================================================

st.markdown(
    """
    <div class="footer">

        ⚖️ <strong>LegalEase</strong>
        &nbsp; · &nbsp;
        AI-Powered Legal Document Workspace
        &nbsp; · &nbsp;
        Built for document drafting assistance

    </div>
    """,
    unsafe_allow_html=True
)